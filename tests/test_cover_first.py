from dataclasses import replace
from io import BytesIO
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from docx import Document
from PIL import Image
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import Paragraph

from src.core.models import CoverImage
from src.exporters.epub import build_epub
from src.exporters.fb2 import build_fb2
from src.exporters import formats
from tests.test_formats_cover import _story


EPUB = '{http://www.idpf.org/2007/ops}'
WORD = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
DRAWING = '{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}'


class CoverFirstTests(unittest.TestCase):
    def test_epub_cover_precedes_toc_annotation_and_chapters(self):
        with ZipFile(BytesIO(build_epub(_story()))) as archive:
            opf = ET.fromstring(archive.read('OEBPS/content.opf'))
            manifest = {node.get('id'): node for node in opf.findall('.//{*}manifest/{*}item')}
            order = [manifest[node.get('idref')].get('href') for node in opf.findall('.//{*}spine/{*}itemref')]
            self.assertEqual(order, ['cover.xhtml', 'nav.xhtml', 'intro.xhtml', 'chapter-0001.xhtml'])
            self.assertEqual(manifest['cover-image'].get('properties'), 'cover-image')
            self.assertEqual(opf.find('.//{*}meta[@name="cover"]').get('content'), 'cover-image')
            self.assertEqual(opf.find('.//{*}guide/{*}reference[@type="cover"]').get('href'), 'cover.xhtml')
            cover = ET.fromstring(archive.read('OEBPS/cover.xhtml'))
            body = cover.find('{*}body')
            self.assertEqual(body.get(EPUB + 'type'), 'cover')
            self.assertEqual(''.join(body.itertext()).strip(), '')
            self.assertIsNone(body.find('.//{*}h1'))
            image = body.find('.//{*}img')
            self.assertEqual(archive.read('OEBPS/' + image.get('src')), _story().cover.data)
            nav = ET.fromstring(archive.read('OEBPS/nav.xhtml'))
            landmark = nav.find(f'.//{{*}}nav[@{EPUB}type="landmarks"]')
            self.assertIsNotNone(landmark)
            self.assertEqual(landmark.find(f'.//{{*}}a[@{EPUB}type="cover"]').get('href'), 'cover.xhtml')

    def test_epub_without_cover_has_no_cover_page_or_landmark(self):
        with ZipFile(BytesIO(build_epub(_story(False)))) as archive:
            self.assertNotIn('OEBPS/cover.xhtml', archive.namelist())
            opf = ET.fromstring(archive.read('OEBPS/content.opf'))
            self.assertIsNone(opf.find('.//{*}item[@id="cover-image"]'))
            self.assertIsNone(opf.find('.//{*}guide'))
            self.assertEqual(opf.find('.//{*}spine/{*}itemref').get('idref'), 'nav')
            nav = ET.fromstring(archive.read('OEBPS/nav.xhtml'))
            self.assertIsNone(nav.find(f'.//{{*}}a[@{EPUB}type="cover"]'))

    def test_fb2_cover_uses_metadata_once_without_a_duplicate_in_the_body(self):
        root = ET.fromstring(build_fb2(_story()))
        cover = root.find('{*}description/{*}title-info/{*}coverpage/{*}image')
        target = cover.get('{http://www.w3.org/1999/xlink}href')[1:]
        self.assertIsNotNone(root.find(f'{{*}}binary[@id="{target}"]'))
        self.assertEqual(len(root.findall('.//{*}coverpage')), 1)
        self.assertIsNone(root.find('{*}body//{*}image'))
        self.assertIsNone(ET.fromstring(build_fb2(_story(False))).find('.//{*}coverpage'))

    def test_docx_cover_is_first_fits_the_page_and_has_no_footer(self):
        for size in ((100, 1200), (1200, 100), (400, 600)):
            with self.subTest(size=size):
                image = BytesIO()
                Image.new('RGB', size, 'navy').save(image, format='PNG')
                story = replace(_story(), cover=CoverImage(image.getvalue(), 'image/png', 'png'))
                document = Document(BytesIO(formats.build_docx(story)))
                section = document.sections[0]
                first = document.paragraphs[0]._p
                extent = first.find('.//' + DRAWING + 'extent')
                self.assertIsNotNone(extent)
                self.assertEqual(document.paragraphs[0].text, '')
                self.assertLessEqual(int(extent.get('cx')), section.page_width - section.left_margin - section.right_margin)
                self.assertLessEqual(int(extent.get('cy')), section.page_height - section.top_margin - section.bottom_margin)
                self.assertAlmostEqual(int(extent.get('cx')) / int(extent.get('cy')), size[0] / size[1], places=4)
                self.assertIsNotNone(document.paragraphs[1]._p.find('.//' + WORD + 'br[@' + WORD + 'type="page"]'))
                self.assertTrue(section.different_first_page_header_footer)
                self.assertFalse(any(p.text for p in section.first_page_footer.paragraphs))

    def test_docx_without_cover_starts_with_text_not_an_empty_page(self):
        document = Document(BytesIO(formats.build_docx(_story(False))))
        self.assertEqual(document.paragraphs[0].text, 'Title')
        self.assertIsNone(document.paragraphs[0]._p.find('.//' + WORD + 'br'))
        self.assertFalse(document.sections[0].different_first_page_header_footer)

    def test_pdf_first_page_contains_only_the_cover_and_no_page_number(self):
        for cover in (True, False):
            with self.subTest(cover=cover):
                images, texts, numbered = [], [], []
                original_image, original_text = Canvas.drawImage, Paragraph.draw

                def draw_image(canvas, *args, **kwargs):
                    images.append(canvas.getPageNumber())
                    return original_image(canvas, *args, **kwargs)

                def draw_text(paragraph):
                    texts.append((paragraph.canv.getPageNumber(), paragraph.getPlainText()))
                    return original_text(paragraph)

                with patch.object(Canvas, 'drawImage', draw_image), \
                        patch.object(Paragraph, 'draw', draw_text), \
                        patch.object(formats, '_draw_pdf_page_number', side_effect=lambda canvas, doc: numbered.append(canvas.getPageNumber())):
                    self.assertTrue(formats.build_pdf(_story(cover)).startswith(b'%PDF'))
                if cover:
                    self.assertEqual(images, [1])
                    self.assertTrue(all(page >= 2 for page, _ in texts))
                    self.assertNotIn(1, numbered)
                    self.assertIn(2, numbered)
                else:
                    self.assertEqual(images, [])
                    self.assertEqual(texts[0][0], 1)
                    self.assertIn(1, numbered)

    def test_txt_is_identical_with_and_without_cover(self):
        self.assertEqual(formats.build_txt(_story()), formats.build_txt(_story(False)))
