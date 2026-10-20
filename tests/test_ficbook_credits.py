from io import BytesIO
import unittest
from unittest.mock import patch
from zipfile import ZipFile
from types import SimpleNamespace
from xml.etree import ElementTree

from bs4 import BeautifulSoup
from docx import Document
from reportlab.platypus import Paragraph

from src.core.models import Chapter, Story
from src.sources.ficbook import FicbookClient
from src.exporters.epub import build_epub
from src.exporters.fb2 import build_fb2
from src.exporters.formats import build_docx, build_pdf, build_txt
from tests.test_ficbook_metadata import FakeStory


class FicbookCreditsTests(unittest.TestCase):
    def test_current_creator_markup_recovers_author_and_ignores_avatar_beta_and_comments(self):
        soup = BeautifulSoup('''<section class="chapter-info"><header><section class="fanfic-hat">
          <div><a href="/authors/1"><img alt="Writer"/></a>
            <div><span><a class="word-break" itemprop="author" href="/authors/1">Writer</a></span>
              <span class="ds-badge ds-badge-regular">автор</span></div></div>
          <div><div><span><a class="word-break" href="/authors/2">Coauthor</a></span>
              <span class="ds-badge ds-badge-regular">соавтор</span></div></div>
          <div><div><span><a class="word-break" href="/authors/3">Beta</a></span>
              <span class="ds-badge ds-badge-regular">бета</span></div></div>
        </section></header></section>
        <div id="comments"><a href="/authors/4">Commenter</a></div>''', 'lxml')
        metadata = FakeStory()
        FicbookClient()._fill_ficbook_author_metadata(metadata, soup)
        self.assertEqual(metadata.getMetadata('author'), 'Writer, Coauthor')
        self.assertEqual(metadata.getMetadata('translator'), '')

    def test_current_translation_keeps_original_author_and_translator_separate(self):
        soup = BeautifulSoup('''<section class="chapter-info"><header><section class="fanfic-hat">
          <div><div><span><a class="word-break" itemprop="author" href="/authors/786035">negative_one</a></span>
            <span class="ds-badge ds-badge-regular">переводчик</span></div></div>
          <div><div><span><a class="word-break" href="/authors/488214">N_a_t_a_l_y</a></span>
            <span class="ds-badge ds-badge-regular">бета</span></div></div>
        </section><div><strong>Автор оригинала:</strong><div>
          <a href="/translations/by_author?author=Rezzekmood">Rezzekmood</a>
        </div></div></header></section>''', 'lxml')
        metadata = FakeStory()
        FicbookClient()._fill_ficbook_author_metadata(metadata, soup)
        self.assertEqual(metadata.getMetadata('author'), 'Rezzekmood')
        self.assertEqual(metadata.getMetadata('translator'), 'negative_one')
        self.assertEqual(metadata.getMetadata('authorUrl'), 'https://ficbook.net/translations/by_author?author=Rezzekmood')

    def test_legacy_avatar_is_skipped_before_text_author_link(self):
        soup = BeautifulSoup('<div id="comments"><a href="/authors/2">Commenter</a></div>'
                             '<a href="/authors/1"><img alt="Writer"/></a><a href="/authors/1">Writer</a>', 'lxml')
        metadata = FakeStory()
        FicbookClient()._fill_ficbook_author_metadata(metadata, soup)
        self.assertEqual(metadata.getMetadata('author'), 'Writer')

    def test_legacy_credits_below_title_header_remain_supported(self):
        soup = BeautifulSoup('''<section class="chapter-info"><header><h1>Title</h1></header>
          <div class="hat-creator-container"><a class="creator-username" href="/authors/1">Writer</a>
            <span class="small-text text-muted">Автор</span></div></section>
          <div id="comments"><a href="/authors/2">Commenter</a></div>''', 'lxml')
        metadata = FakeStory()
        FicbookClient()._fill_ficbook_author_metadata(metadata, soup)
        self.assertEqual(metadata.getMetadata('author'), 'Writer')

    def test_current_badges_are_read_from_work_not_recommendations(self):
        soup = BeautifulSoup('''<aside><div class="ds-label ds-label-rating-NC21">NC-21</div>
          <div class="direction">Гет</div><div class="badge-status-in-progress">В процессе</div></aside>
          <section class="chapter-info"><header><section class="fanfic-badges">
            <div class="direction direction-gen"><span class="hidden-xs">Джен</span></div>
            <div class="ds-label ds-label-rating-R">R</div>
            <div class="ds-label ds-label-status-finished">Завершён</div>
          </section></header></section>''', 'lxml')
        metadata = FakeStory()
        FicbookClient()._fill_optional_ficbook_metadata(SimpleNamespace(story=metadata), soup, get_cover=False)
        self.assertEqual(metadata.getMetadata('classification'), 'Джен')
        self.assertEqual(metadata.getMetadata('rating'), 'R')
        self.assertEqual(metadata.getMetadata('status'), 'Completed')

    def test_annotation_keeps_each_label_with_its_value(self):
        metadata = FakeStory({'title': 'Title', 'author': 'Writer', 'status': 'Completed',
                              'dateStart': '24.05.2019', 'dateFinish': '01.06.2019'})
        html = FicbookClient()._build_annotation_html(metadata, BeautifulSoup('', 'lxml'), 'https://ficbook.net/readfic/1')
        paragraphs = [p.get_text(' ', strip=True) for p in BeautifulSoup(html, 'lxml').find_all('p')]
        for value in ('Автор: Writer', 'Статус: завершён', 'Дата начала: 24.05.2019', 'Дата завершения: 01.06.2019'):
            self.assertIn(value, paragraphs)
        self.assertFalse(any(p.endswith(':') for p in paragraphs))

    def test_author_is_set_in_epub_fb2_docx_and_pdf_metadata(self):
        story = Story('Title', 'MetadataWriter', 'https://ficbook.net/readfic/1', translator='MetadataTranslator')
        with ZipFile(BytesIO(build_epub(story))) as archive:
            root = ElementTree.fromstring(archive.read('OEBPS/content.opf'))
        dc = '{http://purl.org/dc/elements/1.1/}'
        self.assertEqual(root.find('.//' + dc + 'creator').text, story.author)
        self.assertEqual(root.find('.//' + dc + 'contributor').text, story.translator)
        root = ElementTree.fromstring(build_fb2(story))
        ns = '{http://www.gribuser.ru/xml/fictionbook/2.0}'
        self.assertEqual(root.find(f'.//{ns}title-info/{ns}author/{ns}nickname').text, story.author)
        self.assertEqual(root.find(f'.//{ns}title-info/{ns}translator/{ns}nickname').text, story.translator)
        self.assertEqual(Document(BytesIO(build_docx(story))).core_properties.author, story.author)
        self.assertTrue(b'/Author (MetadataWriter)' in build_pdf(story), 'PDF author metadata is missing')

    def test_download_passes_credits_and_complete_annotation_to_exporters(self):
        from fanficfare import adapters
        from src.sources.ficbook import FicbookAccount
        client = FicbookClient()
        url = "https://ficbook.net/readfic/123"
        adapter = adapters.getAdapter(client._build_config(url, FicbookAccount()), url)
        html = '''<section class="chapter-info"><h1>Title</h1></section>
        <div class="hat-creator-container"><a class="creator-username" href="/authors/1">Writer</a><span class="small-text text-muted">Автор</span></div>
        <div class="hat-creator-container"><a class="creator-username" href="/authors/2">Translator</a><span class="small-text text-muted">Переводчик</span></div>
        <ul class="list-of-fanfic-parts"><li class="part"><a href="/readfic/123/1">Chapter</a></li></ul>
        <strong>Размер:</strong><div>12 страниц, 3456 слов</div>'''
        with patch("src.sources.ficbook.adapters.getAdapter", return_value=adapter), patch.object(
            adapter, "get_request", return_value=html
        ), patch.object(client, "_pause_between_requests"), patch.object(client, "_load_chapter_html", return_value="<p>Body</p>"):
            story = client._download_once(url, FicbookAccount(), None, None)
        self.assertEqual(story.author, "Writer")
        self.assertEqual(story.translator, "Translator")
        self.assertTrue(story.annotation_complete)
        self.assertIn("Переводчик: Translator", story.annotation_html)
        self.assertIn("Страниц: 12", story.annotation_html)

    def test_translator_survives_real_fanficfare_metadata(self):
        from fanficfare import adapters
        from src.sources.ficbook import FicbookAccount
        client = FicbookClient()
        url = "https://ficbook.net/readfic/123"
        adapter = adapters.getAdapter(client._build_config(url, FicbookAccount()), url)
        soup = BeautifulSoup('<div class="hat-creator-container"><a class="creator-username" href="/authors/2">Translator</a><span class="small-text text-muted">Переводчик</span></div>', 'lxml')
        client._fill_ficbook_author_metadata(adapter.story, soup)
        self.assertEqual(adapter.story.getMetadata("translator"), "Translator")

    def test_roles_and_multiple_authors(self):
        soup = BeautifulSoup('''<div class="hat-creator-container"><a class="creator-username" href="/authors/1">Beta</a><span class="small-text text-muted">Бета</span></div>
        <div class="hat-creator-container"><a class="creator-username" href="/authors/2">Translator</a><span class="small-text text-muted">Переводчик</span></div>
        <div class="hat-creator-container"><a class="creator-username" href="/authors/3">Writer</a><span class="small-text text-muted">Автор</span></div>
        <div class="hat-creator-container"><a class="creator-username" href="/authors/4">Coauthor</a><span class="small-text text-muted">Соавтор</span></div>''', 'lxml')
        metadata = FakeStory({'author': 'unknown'})
        FicbookClient()._fill_ficbook_author_metadata(metadata, soup)
        self.assertEqual(metadata.getMetadata('author'), 'Writer, Coauthor')
        self.assertEqual(metadata.getMetadata('translator'), 'Translator')

    def test_external_original_author(self):
        soup = BeautifulSoup('''<div class="hat-creator-container"><a class="creator-username" href="/authors/2">Translator</a><span class="small-text text-muted">Переводчик</span></div>
        <strong>Автор оригинала:</strong><div><a href="https://example.org/writer">Original</a></div>''', 'lxml')
        metadata = FakeStory()
        FicbookClient()._fill_ficbook_author_metadata(metadata, soup)
        self.assertEqual(metadata.getMetadata('author'), 'Original')
        self.assertEqual(metadata.getMetadata('translator'), 'Translator')

    def test_missing_original_author_is_not_replaced_by_translator(self):
        soup = BeautifulSoup('<div class="hat-creator-container"><a class="creator-username" href="/authors/2">Translator</a><span class="small-text text-muted">Переводчик</span></div>', 'lxml')
        metadata = FakeStory({'author': 'Translator'})
        FicbookClient()._fill_ficbook_author_metadata(metadata, soup)
        self.assertEqual(metadata.getMetadata('author'), '')
        self.assertEqual(metadata.getMetadata('translator'), 'Translator')

    def test_annotation_escapes_metadata_and_omits_empty_fields(self):
        metadata = FakeStory({'title': 'A < B & C', 'author': '<Writer>', 'pages': '12'})
        html = FicbookClient()._build_annotation_html(metadata, BeautifulSoup('', 'lxml'), 'https://ficbook.net/readfic/1')
        self.assertIn('A &lt; B &amp; C', html)
        self.assertIn('&lt;Writer&gt;', html)
        self.assertNotIn('Переводчик:', html)
        self.assertEqual(html.count('Страниц: 12'), 1)

    def test_all_formats_have_one_ordered_frontmatter(self):
        metadata = FakeStory({'title': 'UniqueTitle', 'author': 'UniqueAuthor', 'translator': 'UniqueTranslator',
                              'pages': '12', 'numWords': '3456', 'description': '<p>UniqueSummary</p>'})
        html = FicbookClient()._build_annotation_html(metadata, BeautifulSoup('', 'lxml'), 'https://ficbook.net/readfic/1')
        story = Story(title='UniqueTitle', author='UniqueAuthor', source_url='https://ficbook.net/readfic/1',
                      translator='UniqueTranslator', annotation_html=html, annotation_complete=True,
                      description='<p>UniqueSummary</p>', page_count='12', word_count='3456', chapters=[Chapter('Chapter', '<p>Body</p>')])
        payloads = {'txt': build_txt(story).decode(), 'docx': '\n'.join(p.text for p in Document(BytesIO(build_docx(story))).paragraphs)}
        payloads['fb2'] = BeautifulSoup(build_fb2(story), 'xml').find('body').get_text(' ')
        with ZipFile(BytesIO(build_epub(story))) as archive:
            payloads['epub'] = BeautifulSoup(archive.read('OEBPS/intro.xhtml'), 'xml').find('body').get_text(' ')
        flow = []
        with patch('src.exporters.formats.SimpleDocTemplate.build', side_effect=lambda items, **kwargs: flow.extend(items)):
            build_pdf(story)
        payloads['pdf'] = '\n'.join(item.getPlainText() for item in flow if isinstance(item, Paragraph))
        for fmt, text in payloads.items():
            with self.subTest(format=fmt):
                for value in ['UniqueTitle', 'UniqueAuthor', 'UniqueTranslator', 'UniqueSummary', 'Страниц: 12', 'Слов: 3456']:
                    self.assertEqual(text.count(value), 1, text)
                self.assertLess(text.index('UniqueTitle'), text.index('UniqueAuthor'))
                self.assertLess(text.index('UniqueAuthor'), text.index('UniqueTranslator'))
                self.assertLess(text.index('UniqueTranslator'), text.index('Страниц: 12'))
                self.assertLess(text.index('Страниц: 12'), text.index('UniqueSummary'))
